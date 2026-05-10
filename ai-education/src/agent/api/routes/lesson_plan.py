# -*- coding: utf-8 -*-
"""
教案生成模块（单文件聚合版）
集成完整智能体 agent_unified，提供路由接口。
"""

import io
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import shutil
from tempfile import NamedTemporaryFile
from typing import Dict, Any, Optional

import markdown
from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from langchain_community.document_loaders import UnstructuredFileLoader
from pydantic import BaseModel
# from weasyprint import HTML

# 导入完整智能体（合并后的文件）
from agent.graphs.lesson_plan.agent_unified import graph, export_to_markdown

# =========================
# 请求/响应模型
# =========================
class LessonRequest(BaseModel):
    grade: str
    subject: str
    course_type: str
    description: Optional[str] = ""

class LessonResponse(BaseModel):
    success: bool
    data: Dict[str, Any]
    markdown: Optional[str] = None

# =========================
# 辅助函数：Markdown 转 HTML（PDF 导出用）
# =========================
def markdown_to_html(md_text: str) -> str:
    body = markdown.markdown(md_text, extensions=['tables', 'fenced_code'])
    return f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <style>
        body {{ font-family: 'Microsoft YaHei', 'SimHei', 'SimSun', 'Noto Sans CJK SC', sans-serif; margin: 2rem; line-height: 1.5; }}
        h1, h2, h3 {{ color: #2c3e50; }}
        table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; }}
        th, td {{ border: 1px solid #ddd; padding: 8px; }}
    </style>
</head>
<body>{body}</body>
</html>
"""

# =========================
# 创建路由器
# =========================
lesson_plan_router = APIRouter()

@lesson_plan_router.post("/generate_lesson", response_model=LessonResponse)
async def generate_lesson(req: LessonRequest):
    try:
        initial_state = {
            "grade": req.grade,
            "subject": req.subject,
            "course_type": req.course_type,
            "messages": [{"role": "user", "content": req.description.strip()}] if req.description and req.description.strip() else [],
        }
        result = await graph.ainvoke(initial_state)
        md = export_to_markdown(result)
        return LessonResponse(success=True, data=result, markdown=md)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成失败: {e}")

@lesson_plan_router.post("/upload_and_generate", response_model=LessonResponse)
async def upload_and_generate(
        grade: str = Form(...),
        subject: str = Form(...),
        course_type: str = Form(...),
        description: str = Form(""),
        file: UploadFile = File(None),
):
    try:
        full_description = description or ""
        if file:
            suffix = os.path.splitext(file.filename or "")[1]
            with NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                shutil.copyfileobj(file.file, tmp)
                tmp_path = tmp.name
            try:
                loader = UnstructuredFileLoader(tmp_path)
                docs = loader.load()
                file_text = "\n".join([d.page_content for d in docs])
                if file_text.strip():
                    full_description += f"\n\n【上传文件内容】\n{file_text}"
            finally:
                os.unlink(tmp_path)

        initial_state = {
            "grade": grade,
            "subject": subject,
            "course_type": course_type,
            "messages": [{"role": "user", "content": full_description.strip()}] if full_description.strip() else [],
        }
        result = await graph.ainvoke(initial_state)
        md = export_to_markdown(result)
        return LessonResponse(success=True, data=result, markdown=md)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成失败: {e}")

@lesson_plan_router.post("/export_pdf")
async def export_pdf(request: Request):
    # try:
    #     markdown_content = (await request.body()).decode("utf-8")
    #     html = markdown_to_html(markdown_content)
    #     pdf_file = HTML(string=html).write_pdf()
    #     return StreamingResponse(
    #         io.BytesIO(pdf_file),
    #         media_type="application/pdf",
    #         headers={"Content-Disposition": "attachment; filename=lesson_plan.pdf"},
    #     )
    # except Exception as e:
    #     raise HTTPException(status_code=500, detail=f"PDF生成失败: {e}")
    raise HTTPException(status_code=501, detail="PDF功能正在配置中，请稍后再试")

@lesson_plan_router.post("/export_docx")
async def export_docx(request: Request):
    try:
        markdown_content = (await request.body()).decode("utf-8")
        doc = Document()
        for line in markdown_content.split("\n"):
            line = line.rstrip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line)
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=lesson_plan.docx"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Word生成失败: {e}")